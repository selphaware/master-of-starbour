"""Generate docs/DESIGN.md, docs/DECISIONS.md and docs/config_keys.json from the
design workbook docs/master_of_starbour_blueprint.xlsx.

The workbook is the source of truth (Project instructions). This script makes the
Markdown copies that Claude Code reads, so the two can never drift silently:
re-run it after every workbook change and commit the result.

Usage (from the repo root, with the MOS interpreter):

    E:\\CODE\\master-of-starbour\\MOS\\Scripts\\python.exe tools\\blueprint_to_design.py
    E:\\CODE\\master-of-starbour\\MOS\\Scripts\\python.exe tools\\blueprint_to_design.py --check

--check writes nothing; it exits 1 if any generated file is out of date.

What it does, tab by tab:
  * every tab except '17 Decisions' -> a section of docs/DESIGN.md, every row kept,
    every ID kept, tables rendered as Markdown tables;
  * '17 Decisions' -> docs/DECISIONS.md (the hand-maintained kit decisions below the
    marker line are preserved);
  * '15 Config Keys' -> docs/config_keys.json, the manifest that
    tests/test_config_matches_blueprint.py checks controls_config.ini against.

Dependencies: openpyxl only (dev dependency).
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("HARNESS ERROR: openpyxl is not installed in this interpreter "
             "(pip install openpyxl). This is not a deliverable problem.")

REPO = Path(__file__).resolve().parents[1]
WORKBOOK = REPO / "docs" / "master_of_starbour_blueprint.xlsx"
DESIGN_MD = REPO / "docs" / "DESIGN.md"
DECISIONS_MD = REPO / "docs" / "DECISIONS.md"
CONFIG_JSON = REPO / "docs" / "config_keys.json"

KIT_MARKER = "<!-- KIT DECISIONS — everything below this line is hand-maintained; the generator keeps it. -->"
ID_RE = re.compile(r"^[A-Z]{1,4}-\d+[a-z]?$")


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.4g}"
    return str(value)


def md_escape(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", "<br>")


def row_values(row) -> list[str]:
    cells = [cell_text(v) for v in row]
    while cells and cells[-1] == "":
        cells.pop()
    return cells


def is_heading(cells: list[str]) -> bool:
    if len([c for c in cells if c]) != 1:
        return False
    text = next(c for c in cells if c)
    return len(text) <= 40 and not text.endswith(".")


HEADER_WORDS = {"ID", "Status", "Origin", "Rows", "Text", "note", "colour", "Tab",
                "Row", "distance", "warp factor", "Example star", "reach (sectors)"}


def looks_like_header(cells: list[str]) -> bool:
    return any(c in HEADER_WORDS for c in cells)


def render_table(block: list[list[str]]) -> list[str]:
    if not looks_like_header(block[0]):
        # A key/value or calculator block with no header row: render as a list so
        # no data row is mistaken for a header.
        out = []
        for r in block:
            parts = [md_escape(c) for c in r if c]
            out.append(f"- **{parts[0]}** — " + " · ".join(parts[1:]) if len(parts) > 1 else f"- {parts[0]}")
        out.append("")
        return out
    width = max(len(r) for r in block)
    rows = [r + [""] * (width - len(r)) for r in block]
    header, body = rows[0], rows[1:]
    out = ["| " + " | ".join(md_escape(c) for c in header) + " |",
           "|" + "---|" * width]
    for r in body:
        out.append("| " + " | ".join(md_escape(c) for c in r) + " |")
    out.append("")
    return out


def render_tab(ws) -> list[str]:
    lines: list[str] = []
    rows = [row_values(r) for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(r)]
    if not rows:
        return lines
    title = rows[0][0]
    lines.append(f"## {ws.title}")
    lines.append("")
    lines.append(f"**{title}**")
    lines.append("")
    body = rows[1:]
    if body and len([c for c in body[0] if c]) == 1:
        lines.append(f"*{body[0][0]}*")
        lines.append("")
        body = body[1:]
    table: list[list[str]] = []
    for cells in body:
        nonempty = [c for c in cells if c]
        if nonempty == ["↩ Index"]:
            continue
        if len(nonempty) >= 2:
            table.append(cells)
            continue
        if table:
            lines.extend(render_table(table))
            table = []
        if is_heading(cells):
            lines.append(f"### {nonempty[0]}")
        else:
            lines.append(nonempty[0])
        lines.append("")
    if table:
        lines.extend(render_table(table))
    return lines


def build_design(wb) -> str:
    index = wb["0 Index"]
    version_line = cell_text(index["A2"].value)
    out = [
        "# Master of Starbour — DESIGN.md",
        "",
        "**Generated from `docs/master_of_starbour_blueprint.xlsx` by "
        "`tools/blueprint_to_design.py`. Do not edit by hand — edit the workbook and "
        "regenerate.**",
        "",
        f"Workbook version: {version_line}",
        "",
        "How to read this file:",
        "",
        "- Every row keeps its stable ID (OV-07, CMD-06, HZ-02, CFG-17 …). Cite IDs in "
        "every spec, test, gate and decision.",
        "- **Agreed** = Usman's decision, ready for a spec. **Proposed** = Claude's "
        "placeholder (a number, a formula, a default) awaiting playtest or Usman's yes — "
        "never treat it as decided. **Draft** = reworded from his blueprint, not yet "
        "confirmed. **Needs detail** / **Not started** = open.",
        "- Every number is a placeholder read from `controls_config.ini` and tuned by "
        "playtest. There are no fixed turn-count targets.",
        "- The decision log (tab 17) is in `docs/DECISIONS.md`; the config-key manifest "
        "(tab 15) is also in `docs/config_keys.json`.",
        "",
        "Contents:",
        "",
    ]
    tabs = [ws for ws in wb.worksheets if ws.title != "17 Decisions"]
    for ws in tabs:
        anchor = re.sub(r"[^a-z0-9]+", "-", ws.title.lower()).strip("-")
        out.append(f"- [{ws.title}](#{anchor})")
    out.append("")
    for ws in tabs:
        out.append("---")
        out.append("")
        out.extend(render_tab(ws))
    return "\n".join(out).rstrip() + "\n"


def build_decisions(wb, existing: str | None) -> str:
    ws = wb["17 Decisions"]
    rows = [row_values(r) for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(r)]
    header_idx = next(i for i, r in enumerate(rows) if r and r[0] == "ID")
    table = rows[header_idx:]
    out = [
        "# Master of Starbour — DECISIONS.md",
        "",
        "Settled decisions. D-01…D-53 are generated from tab 17 of "
        "`docs/master_of_starbour_blueprint.xlsx` by `tools/blueprint_to_design.py` — do "
        "not edit them here. Kit decisions (D-54 onwards) are hand-maintained below the "
        "marker and are kept by the generator.",
        "",
        "## From the workbook (tab 17)",
        "",
    ]
    out.extend(render_table(table))
    out.append(KIT_MARKER)
    out.append("")
    if existing and KIT_MARKER in existing:
        out.append(existing.split(KIT_MARKER, 1)[1].lstrip("\n"))
    else:
        out.append("## Kit decisions (hand-maintained)")
        out.append("")
        out.append("| ID | Date | Decision | Rationale / source | Ref | Status |")
        out.append("|---|---|---|---|---|---|")
        out.append("")
    return "\n".join(out).rstrip() + "\n"


def build_config_manifest(wb) -> str:
    ws = wb["15 Config Keys"]
    rows = [row_values(r) for r in ws.iter_rows(values_only=True)]
    header_idx = next(i for i, r in enumerate(rows) if r and r[0] == "ID")
    header = rows[header_idx]
    keys = []
    for r in rows[header_idx + 1:]:
        if not r or not ID_RE.match(r[0]):
            continue
        r = r + [""] * (len(header) - len(r))
        rec = dict(zip(header, r))
        section = rec["Section"].strip("[]")
        names = [k.strip() for k in rec["Key"].split(",")]
        keys.append({
            "id": rec["ID"],
            "section": section,
            "keys": names,
            "type": rec["Type"],
            "default": rec["Default"],
            "unit": rec["Unit / range"],
            "meaning": rec["Meaning"],
            "used_by": rec["Used by"],
            "status": rec["Status"],
        })
    return json.dumps({"source": WORKBOOK.name, "keys": keys}, indent=2, ensure_ascii=False) + "\n"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true", help="verify generated files are current; write nothing")
    args = ap.parse_args()
    if not WORKBOOK.exists():
        print(f"HARNESS ERROR: workbook not found at {WORKBOOK}")
        return 3
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    existing_decisions = DECISIONS_MD.read_text(encoding="utf-8") if DECISIONS_MD.exists() else None
    outputs = {
        DESIGN_MD: build_design(wb),
        DECISIONS_MD: build_decisions(wb, existing_decisions),
        CONFIG_JSON: build_config_manifest(wb),
    }
    stale = []
    for path, content in outputs.items():
        current = path.read_text(encoding="utf-8") if path.exists() else None
        if current != content:
            stale.append(path)
            if not args.check:
                path.write_text(content, encoding="utf-8", newline="\n")
                print(f"wrote {path.relative_to(REPO)} ({len(content):,} chars)")
    if args.check:
        if stale:
            print("STALE: " + ", ".join(str(p.relative_to(REPO)) for p in stale)
                  + " — run tools/blueprint_to_design.py and commit")
            return 1
        print("docs are current with the workbook")
    elif not stale:
        print("nothing to do — docs already current")
    return 0


if __name__ == "__main__":
    sys.exit(main())
